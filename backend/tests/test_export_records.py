from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker

from app.db.models.cccd_record import CCCDRecord
from app.db.models.correction import Correction
from app.db.models.scan import ScanStatus
from app.db.repositories.scan_repository import ScanRepository

OCR_DATA = {
    "id_number": "001202000123",
    "full_name": "NGUYEN VAN A",
    "dob": "01/01/2002",
    "gender": "Nam",
    "nationality": "Việt Nam",
    "origin": "Hà Nội",
    "residence": "Hà Nội",
    "expiry_date": "01/01/2027",
}


def _headers(client: TestClient, email: str) -> dict[str, str]:
    client.post(
        "/api/auth/register",
        json={"email": email, "password": "secret123", "full_name": "Tester"},
    )
    token = client.post(
        "/api/auth/login", json={"email": email, "password": "secret123"}
    ).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _me(client: TestClient, headers: dict[str, str]) -> str:
    return client.get("/api/auth/me", headers=headers).json()["id"]


def _create_scan(client: TestClient, factory: async_sessionmaker[AsyncSession], user_id: str | None) -> str:
    async def _create() -> str:
        async with factory() as session:
            scan = await ScanRepository(session).create_scan(
                user_id=user_id, request_id=None, filename="a.jpg", mime_type="image/jpeg",
                size_bytes=1, image_hash="h", ocr_engine="mock", status=ScanStatus.SUCCESS,
                raw_data=OCR_DATA, confidences={"full_name": 0.6}, error_code=None,
                error_message=None, duration_ms=1,
            )
            await session.commit()
            return scan.id

    return client.portal.call(_create)


def _count(client: TestClient, factory: async_sessionmaker[AsyncSession], model) -> int:  # noqa: ANN001
    async def _run() -> int:
        async with factory() as session:
            return (await session.execute(select(func.count()).select_from(model))).scalar_one()

    return client.portal.call(_run)


def test_reexporting_same_scan_updates_record_instead_of_duplicating(
    db_client: TestClient, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    headers = _headers(db_client, "export@example.com")
    scan_id = _create_scan(db_client, session_factory, _me(db_client, headers))

    first = {**OCR_DATA, "full_name": "NGUYEN VAN B"}
    second = {**OCR_DATA, "full_name": "NGUYEN VAN C", "expiry_date": None}
    for data in (first, second):
        res = db_client.post("/api/export", json={"data": data, "scan_id": scan_id}, headers=headers)
        assert res.status_code == 200

    records = db_client.get("/api/records", headers=headers).json()
    assert len(records) == 1
    assert records[0]["full_name"] == "NGUYEN VAN C"
    assert records[0]["expiry_date"] is None
    assert _count(db_client, session_factory, Correction) == 2


def test_export_rejects_scan_of_another_user(
    db_client: TestClient, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    owner = _headers(db_client, "owner@example.com")
    scan_id = _create_scan(db_client, session_factory, _me(db_client, owner))
    other = _headers(db_client, "other@example.com")

    res = db_client.post("/api/export", json={"data": OCR_DATA, "scan_id": scan_id}, headers=other)
    assert res.status_code == 403
    assert _count(db_client, session_factory, Correction) == 0
    assert _count(db_client, session_factory, CCCDRecord) == 0


def test_excel_cells_do_not_become_formulas(db_client: TestClient) -> None:
    data = {**OCR_DATA, "residence": '=HYPERLINK("http://evil","x")'}
    res = db_client.post("/api/export", json={"data": data})
    assert res.status_code == 200
    sheet = load_workbook(io.BytesIO(res.content)).active
    cell = sheet.cell(row=2, column=7)
    assert cell.data_type == "s"
    assert cell.value.startswith("'=")


@pytest.mark.parametrize("bad", [{"gender": "X"}, {"id_number": "00120200012a"}])
def test_record_create_rejects_invalid_fields_with_422(db_client: TestClient, bad: dict) -> None:
    headers = _headers(db_client, "valid@example.com")
    res = db_client.post("/api/records", json={**OCR_DATA, **bad}, headers=headers)
    assert res.status_code == 422


def test_record_update_can_fix_id_number_and_missing_record_is_404(db_client: TestClient) -> None:
    headers = _headers(db_client, "patch@example.com")
    record_id = db_client.post("/api/records", json=OCR_DATA, headers=headers).json()["id"]

    res = db_client.patch(f"/api/records/{record_id}", json={"id_number": "079202000999"}, headers=headers)
    assert res.status_code == 200
    assert res.json()["id_number"] == "079202000999"

    assert db_client.patch("/api/records/missing", json={}, headers=headers).status_code == 404
    assert db_client.get("/api/scans/missing", headers=headers).status_code == 404


LIST_ROW = {
    "full_name": "NGUYEN VAN A",
    "id_number": "001202000123",
    "dob": "01/01/2002",
    "gender": "Nam",
    "residence": "Hà Nội",
}


def test_export_list_writes_one_row_per_person(db_client: TestClient) -> None:
    second = {**LIST_ROW, "full_name": "TRAN THI B", "id_number": "079305000456", "dob": "02/03/2005", "gender": "Nữ"}
    res = db_client.post("/api/export-list", json={"rows": [LIST_ROW, second]})
    assert res.status_code == 200
    assert "DanhSach_CCCD_" in res.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(res.content)).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("STT", "Họ và tên", "Số CCCD", "Ngày sinh", "Giới tính", "Nơi thường trú")
    assert rows[1] == (1, "NGUYEN VAN A", "001202000123", "01/01/2002", "Nam", "Hà Nội")
    assert rows[2][:3] == (2, "TRAN THI B", "079305000456")
    assert sheet.cell(row=2, column=3).number_format == "@"


@pytest.mark.parametrize(
    "bad",
    [
        {"full_name": "(không nhận dạng được)"},
        {"id_number": "000000000000"},
        {"gender": "Nữ"},
        {"dob": "01/01/1999"},
        {"residence": ""},
        {"full_name": "NGUYEN VAN A1"},
        {"id_number": "001299000123", "dob": "01/01/2099"},
    ],
)
def test_export_list_rejects_unreviewed_or_inconsistent_rows(db_client: TestClient, bad: dict) -> None:
    res = db_client.post("/api/export-list", json={"rows": [LIST_ROW, {**LIST_ROW, **bad}]})
    assert res.status_code == 422
    assert res.json()["details"][0]["loc"][:3] == ["body", "rows", 1]


def test_export_list_logs_corrections_only_for_fields_the_scan_read(
    db_client: TestClient, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    scan_id = _create_scan(db_client, session_factory, None)
    row = {**LIST_ROW, "full_name": "NGUYỄN VĂN A", "residence": "Hà Nội mới", "scan_ids": [scan_id]}
    assert db_client.post("/api/export-list", json={"rows": [row]}).status_code == 200

    async def _fields() -> list[str]:
        async with session_factory() as session:
            return list((await session.execute(select(Correction.field_name))).scalars())

    assert db_client.portal.call(_fields) == ["full_name"]


def test_export_list_still_exports_but_never_writes_to_another_users_scan(
    db_client: TestClient, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    owner = _headers(db_client, "list-owner@example.com")
    scan_id = _create_scan(db_client, session_factory, _me(db_client, owner))
    row = {**LIST_ROW, "full_name": "TÊN KHÁC", "scan_ids": [scan_id]}
    res = db_client.post("/api/export-list", json={"rows": [row]})
    assert res.status_code == 200
    assert _count(db_client, session_factory, Correction) == 0


def test_anonymous_scan_is_not_readable_by_other_accounts(
    db_client: TestClient, session_factory: async_sessionmaker[AsyncSession]
) -> None:
    scan_id = _create_scan(db_client, session_factory, None)
    headers = _headers(db_client, "curious@example.com")
    assert db_client.get(f"/api/scans/{scan_id}", headers=headers).status_code == 403
